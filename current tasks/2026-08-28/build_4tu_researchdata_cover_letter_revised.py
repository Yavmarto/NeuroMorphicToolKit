from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).with_name(
    "Yoshi Martodihardjo-Bink - Cover Letter - TU Delft 4TU ResearchData - Revised.docx"
)


def set_run_font(run, size=11, bold=None, color="000000"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def format_paragraph(paragraph, before=0, after=8, line=1.10):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    p_pr.append(borders)


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, after=8)
    set_run_font(paragraph.add_run(text))


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.74)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.90)
    section.right_margin = Inches(0.90)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    name = doc.add_paragraph()
    format_paragraph(name, after=1)
    set_run_font(name.add_run("Yoshi Martodihardjo-Bink"), size=18, bold=True, color="0B2545")

    contact = doc.add_paragraph()
    format_paragraph(contact, after=7)
    set_run_font(
        contact.add_run(
            "Etten-Leur, Netherlands  |  +31 6 55585355  |  yoshi.martodihardjo@gmail.com"
        ),
        size=9.5,
        color="4F5B66",
    )
    add_bottom_rule(contact)

    date = doc.add_paragraph()
    format_paragraph(date, after=10)
    set_run_font(date.add_run("28 August 2026"), size=10.5)

    recipient = doc.add_paragraph()
    format_paragraph(recipient, after=9)
    set_run_font(
        recipient.add_run(
            "Curtis Sharma\n4TU.ResearchData / TU Delft Library\nDelft, The Netherlands"
        ),
        size=10.5,
    )

    subject = doc.add_paragraph()
    format_paragraph(subject, after=9)
    set_run_font(
        subject.add_run("Re: Software Engineer (Open Source — 4TU.ResearchData)"),
        size=10.5,
        bold=True,
        color="1F4D78",
    )

    greeting = doc.add_paragraph()
    format_paragraph(greeting, after=8)
    set_run_font(greeting.add_run("Dear Curtis Sharma,"))

    add_body(
        doc,
        "I am applying for the Software Engineer role because 4TU.ResearchData’s mission is the problem I have been working on from the researcher’s side: making research software understandable, reproducible, and useful beyond the person who first built it. I am a Dutch software engineer with eight years of professional delivery experience, currently a lead cross-platform engineer and an MSc Applied Artificial Intelligence student."
    )
    add_body(
        doc,
        "My strongest evidence is NeuroMorphicToolKit (NMTK), the public AGPL-licensed research-software platform I have architected and built for neuromorphic engineering. It is precisely the kind of software that should be discoverable beyond its source repository. Its GitHub project can remain the place where development happens, while an RSD entry could make its purpose, people, licence, documentation, associated research, and reuse visible to the research community; a versioned 4TU.ResearchData deposit could provide a stable, citable release snapshot."
    )
    add_body(
        doc,
        "That is not a hypothetical I treat casually. NMTK already has public source, a clear licence, architecture and contribution documentation, and named copyright and contributor information. To make it citation-ready, I would need to complete the remaining metadata and release work deliberately: establish a citation record, connect the people and research context, archive an exact release, and make the resulting identifier and documentation part of the project’s normal workflow. I understand why a directory and repository add value here: they do not duplicate GitHub; they make software findable, attributable, assessable, and citable."
    )
    add_body(
        doc,
        "Building NMTK has also made the practical meaning of reproducibility concrete. It takes a spiking neural network from a controlled-English or visual specification through a hardware-independent representation, simulation, benchmarking, and deployment. I designed complete workspaces around the model, weights, target configuration, and benchmark results, and made support states explicit—works, needs hardware, or not implemented—alongside faithful, approximate, and unsupported exports. A future user should be able to see what ran, on which target, with what limitations, rather than infer it from incomplete context."
    )
    add_body(
        doc,
        "This is why I fit the role: I can contribute as both a Python engineer and a demanding future user of Djehuty and the Research Software Directory. My work includes Python, FastAPI, PyTorch, GNU/Linux, Git, Docker/Podman, pytest, Ruff, GitHub Actions, and JupyterLab, but equally important is my experience translating complex requirements into a workflow researchers can use. Alongside this, I lead cross-platform product delivery from architecture to release, including UX/UI, testing, code review, and mentoring. I would be excited to help build sustainable open-source infrastructure that makes it easier for the next NMTK-like project to be found, reused, and credited properly."
    )

    closing = doc.add_paragraph()
    format_paragraph(closing, after=12)
    set_run_font(closing.add_run("Yours sincerely,"))
    signature = doc.add_paragraph()
    format_paragraph(signature, after=0)
    set_run_font(signature.add_run("Yoshi Martodihardjo-Bink"), bold=True, color="0B2545")

    doc.core_properties.title = "Cover Letter — TU Delft 4TU.ResearchData"
    doc.core_properties.author = "Yoshi Martodihardjo-Bink"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
